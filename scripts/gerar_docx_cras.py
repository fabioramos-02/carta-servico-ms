"""Gera .docx da carta 'Consultar endereços dos CRAS' em Arial 12 com hiperlinks."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "exemplos" / "saida" / "consultar-enderecos-cras" / "consultar-enderecos-cras.docx"

FONT = "Arial"
SIZE = 12

doc = Document()

style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(SIZE)
rpr = style.element.get_or_add_rPr()
rfonts = rpr.find(qn("w:rFonts"))
if rfonts is None:
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rfonts.set(qn(attr), FONT)


def set_run_font(run, bold=False, size=SIZE):
    run.font.name = FONT
    run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)
    run.bold = bold


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(SIZE * 2))
    rPr.append(sz)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    size = {1: 18, 2: 14, 3: 12}[level]
    set_run_font(run, bold=True, size=size)
    return p


def add_para(text=""):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_para_rich(segments):
    p = doc.add_paragraph()
    for seg in segments:
        if seg[0] == "text":
            _, txt, bold = seg
            run = p.add_run(txt)
            set_run_font(run, bold=bold)
        elif seg[0] == "link":
            _, url, label = seg
            add_hyperlink(p, url, label)
    return p


def add_bullet(segments, numbered=False, sub=False):
    style_name = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style_name)
    if sub:
        p.paragraph_format.left_indent = Pt(36)
    for seg in segments:
        if seg[0] == "text":
            _, txt, bold = seg
            run = p.add_run(txt)
            set_run_font(run, bold=bold)
        elif seg[0] == "link":
            _, url, label = seg
            add_hyperlink(p, url, label)
    return p


PLAY = "https://play.google.com/store/apps/details?id=br.gov.ms.msapp"
APPSTORE = "https://apps.apple.com/br/app/ms-digital/id1482970942"

# =========== CONTEÚDO ===========

add_heading("Consultar endereços dos CRAS", 1)

add_heading("O QUE É?", 2)
add_para(
    "Serviço on-line em que você consulta o endereço, telefone e e-mail dos Centros "
    "de Referência de Assistência Social (CRAS) do Estado de Mato Grosso do Sul. "
    "Você faz a busca pelo aplicativo MS Digital, escolhendo o município e a unidade desejada."
)

add_heading("Exigências", 2)
add_bullet([
    ("text", "Aplicativo MS Digital instalado (", False),
    ("link", PLAY, "Google Play"),
    ("text", " / ", False),
    ("link", APPSTORE, "App Store"),
    ("text", ")", False),
], numbered=True)
add_bullet([("text", "Conta gov.br autenticada ou CPF e senha", False)], numbered=True)

add_heading("Quem pode utilizar?", 2)
add_para("Qualquer cidadão que precise localizar um CRAS em Mato Grosso do Sul.")

add_heading("Prazo", 2)
add_para("Imediato.")

add_heading("Custos", 2)
add_para("Sem custos.")

add_heading("Etapas", 2)

add_para_rich([("text", "Pelo aplicativo MS Digital:", True)])
add_bullet([
    ("text", "Baixe ", True),
    ("text", "o aplicativo MS Digital na ", False),
    ("link", PLAY, "Google Play"),
    ("text", " (Android) ou ", False),
    ("link", APPSTORE, "App Store"),
    ("text", " (iPhone).", False),
], numbered=True)
add_bullet([
    ("text", "Faça login ", True),
    ("text", "com sua conta gov.br ou com CPF e senha.", False),
], numbered=True)
add_bullet([
    ("text", "Acesse ", True),
    ("text", "a categoria Assistência Social.", False),
], numbered=True)
add_bullet([
    ("text", "Selecione ", True),
    ("text", "o serviço Endereços dos CRAS.", False),
], numbered=True)
add_bullet([
    ("text", "Escolha ", True),
    ("text", "o município desejado.", False),
], numbered=True)
add_bullet([
    ("text", "Selecione ", True),
    ("text", "o CRAS. A tela mostra telefone, e-mail e endereço da unidade.", False),
], numbered=True)
add_bullet([
    ("text", "Toque ", True),
    ("text", "no telefone para ligar ou no endereço para abrir a localização no mapa.", False),
], numbered=True)

add_heading("Outras Informações", 2)

add_para_rich([("text", "Consulta alternativa pela internet:", True)])
add_bullet([
    ("link", "https://www.sead.ms.gov.br/fale-conosco/unidades/unidades-regionais/", "Unidades Regionais da SEAD"),
    ("text", " — lista completa das unidades e contatos no site oficial", False),
])

add_para_rich([
    ("text", "Órgão responsável: ", True),
    ("text", "Secretaria de Estado de Assistência Social — SEAD", False),
])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(f"OK -> {OUT}")
